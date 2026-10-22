"""Build excel/olist_analysis.xlsx from dashboard_data/ (run build_dashboard_data.py first).

Sheets: README, KPI_Monthly (with a native line chart), Top_Categories, State_Delivery,
Segments, and Orders_Data (delivered orders) - the sheet you use to insert your own PivotTable.

Usage:  python scripts/make_excel.py
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
D = ROOT / "dashboard_data"
OUT = ROOT / "excel" / "olist_analysis.xlsx"
HEAD_FILL = PatternFill("solid", fgColor="2A6FDB")


def write_df(ws, df: pd.DataFrame, start_row: int = 1) -> None:
    for j, col in enumerate(df.columns, 1):
        c = ws.cell(row=start_row, column=j, value=col)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), HEAD_FILL
        c.alignment = Alignment(horizontal="center")
    for i, row in enumerate(df.itertuples(index=False), start_row + 1):
        for j, v in enumerate(row, 1):
            if pd.isna(v):
                v = None
            elif isinstance(v, pd.Timestamp):
                v = v.to_pydatetime()
            ws.cell(row=i, column=j, value=v)
    for j, col in enumerate(df.columns, 1):
        ws.column_dimensions[get_column_letter(j)].width = max(12, min(30, len(str(col)) + 4))
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)


def main() -> None:
    OUT.parent.mkdir(exist_ok=True)
    kpi = pd.read_csv(D / "kpi_monthly.csv", parse_dates=["month"])
    cat = pd.read_csv(D / "category_revenue.csv").head(15)
    st = pd.read_csv(D / "state_delivery.csv")
    seg = pd.read_csv(D / "segment_summary.csv")
    fo = pd.read_csv(D / "fact_orders.csv", parse_dates=["purchase_date", "month"])
    fo = fo[fo["is_delivered"] == 1][
        ["order_id", "customer_state", "purchase_date", "month", "items_revenue", "freight",
         "review_score", "delivery_days", "is_late"]
    ].copy()
    fo["month"] = fo["month"].dt.strftime("%Y-%m")
    fo["delivery_status"] = fo["is_late"].map({1.0: "Late", 0.0: "On time"})
    fo = fo.drop(columns="is_late")

    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    for i, line in enumerate([
        "Olist e-commerce analysis - Excel workbook",
        "Revenue = item price excluding freight, delivered orders, Jan 2017 - Aug 2018 (KPI sheets).",
        "Orders_Data = one row per delivered order. Build your PivotTable from it:",
        "  Insert > PivotTable > select Orders_Data table > Rows: month, Columns: delivery_status, Values: Average of review_score.",
        "  Second pivot: Rows: customer_state, Values: Sum of items_revenue and Average of delivery_days; add a Slicer on month.",
    ], 1):
        ws.cell(row=i, column=1, value=line).font = Font(bold=(i == 1), size=14 if i == 1 else 11)
    ws.column_dimensions["A"].width = 120

    ws = wb.create_sheet("KPI_Monthly")
    k = kpi.copy()
    k["month"] = k["month"].dt.strftime("%Y-%m")
    write_df(ws, k)
    n = len(k) + 1
    ch = LineChart()
    ch.title, ch.height, ch.width = "Monthly revenue", 8, 18
    ch.add_data(Reference(ws, min_col=2, min_row=1, max_row=n), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=2, max_row=n))
    ws.add_chart(ch, "K2")

    ws = wb.create_sheet("Top_Categories")
    write_df(ws, cat)
    bc = BarChart()
    bc.type, bc.title, bc.height, bc.width = "bar", "Top categories by revenue", 9, 18
    bc.add_data(Reference(ws, min_col=2, min_row=1, max_row=len(cat) + 1), titles_from_data=True)
    bc.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(cat) + 1))
    ws.add_chart(bc, "H2")

    write_df(wb.create_sheet("State_Delivery"), st)
    write_df(wb.create_sheet("Segments"), seg)
    write_df(wb.create_sheet("Orders_Data"), fo)

    wb.save(OUT)
    print("saved", OUT, f"({len(fo):,} delivered orders)")


if __name__ == "__main__":
    main()
